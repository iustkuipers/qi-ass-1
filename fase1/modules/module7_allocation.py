"""Module 7 - the capital allocation and the hand-in sheet.

Two separate decisions, in the order week 2 takes them:

  1. *which* risky portfolio to hold - settled in step 6, and independent of
     risk aversion because every mean-variance investor holds the same
     tangency portfolio (the two-fund separation result);
  2. *how much* of the budget to put in it, which is where risk aversion
     enters:

         y* = (E[r_p] - rf) / (A * sigma_p^2)

     with the rest at the risk-free rate. Borrowing is not allowed, so y* is
     capped at 1.

Then the practical part: EUR 10 million buys a whole number of shares, so
every position is rounded down and the cash left over joins the risk-free
holding. Rounding down rather than to nearest guarantees the budget is never
exceeded, and the remainder is never wasted because it earns the risk-free
rate rather than sitting idle.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from config import (BUDGET, RF_ANNUAL, SHEET_FIRST_FIRM_ROW,  # noqa: F401
                    SHEET_LAST_FIRM_ROW, SHEET_RF_ROW)


def optimal_risky_share(excess_return: float, volatility: float,
                        risk_aversion: float,
                        allow_borrowing: bool = False) -> dict:
    """y* = (E[r] - rf) / (A sigma^2), optionally capped at 1."""
    variance = volatility ** 2
    raw = excess_return / (risk_aversion * variance)
    y = raw if allow_borrowing else min(raw, 1.0)
    return {
        "risk_aversion": risk_aversion,
        "y_raw": raw,
        "y": y,
        "binding": (not allow_borrowing) and raw > 1.0,
        "risky_amount": y * BUDGET,
        "riskfree_amount": (1 - y) * BUDGET,
    }


def allocate(weights: pd.Series, prices: pd.Series, y: float,
             budget: float = BUDGET) -> tuple[pd.DataFrame, dict]:
    """Turn portfolio weights into a whole number of shares.

    Every position is rounded *down*. The alternative, rounding to nearest,
    can spend more than the budget, and the assignment gives us exactly EUR 10
    million. What is left over goes to the risk-free asset, so nothing is lost:
    it earns 1% a year rather than nothing.
    """
    target_value = weights * y * budget
    shares = np.floor(target_value / prices.reindex(weights.index)).astype("int64")
    shares = shares.clip(lower=0)

    value = shares * prices.reindex(weights.index)
    invested = float(value.sum())
    riskfree = budget - invested

    table = pd.DataFrame({
        "weight_target": weights,
        "price": prices.reindex(weights.index),
        "target_value": target_value,
        "shares": shares,
        "value": value,
        "weight_actual": value / budget,
    })

    info = {
        "invested": invested,
        "riskfree": riskfree,
        "total": invested + riskfree,
        "n_positions": int((shares > 0).sum()),
        "n_dropped_by_rounding": int(((target_value > 0) & (shares == 0)).sum()),
        "rounding_cash": float(y * budget - invested),
        "deliberate_riskfree": float((1 - y) * budget),
        "largest_position": float(value.max()),
        "smallest_position": float(value[value > 0].min()) if (value > 0).any() else 0.0,
    }
    return table, info


def realised_portfolio_stats(table: pd.DataFrame, mu_excess: pd.Series,
                             sigma: pd.DataFrame, budget: float = BUDGET) -> dict:
    """What the *rounded* portfolio actually is, not what we asked for.

    Rounding moves the weights a little, so the expected return and volatility
    are recomputed on the shares we really hold.
    """
    held = table.loc[table["shares"] > 0]
    if held.empty:
        return {}
    w = held["value"] / held["value"].sum()          # weights within the risky book
    assets = list(w.index)
    mu = mu_excess.reindex(assets).to_numpy()
    S = sigma.loc[assets, assets].to_numpy()
    wv = w.to_numpy()

    risky_excess = float(wv @ mu)
    risky_vol = float(np.sqrt(wv @ S @ wv))

    y_actual = float(held["value"].sum() / budget)
    total_excess = y_actual * risky_excess
    total_vol = y_actual * risky_vol

    return {
        "y_actual": y_actual,
        "risky_excess_return": risky_excess,
        "risky_volatility": risky_vol,
        "risky_sharpe": risky_excess / risky_vol,
        "total_excess_return": total_excess,
        "total_expected_return": total_excess + RF_ANNUAL,
        "total_volatility": total_vol,
        "tracking_to_target": float(np.abs(
            held["weight_actual"] - held["weight_target"] * y_actual).max()),
    }


def fill_hand_in_sheet(table: pd.DataFrame, universe: pd.DataFrame,
                       riskfree_amount: float, source: Path, target: Path,
                       shares_column: str = "H",
                       group: dict | None = None) -> dict:
    """Write the share counts into the hand-in workbook.

    The sheet has 505 firm rows for 504 firms - HST is listed twice - so the
    shares go on the first of the two rows and the repeat gets zero. The
    risk-free line has a price of 1, so the number written there is simply the
    euro amount.
    """
    import openpyxl

    wb = openpyxl.load_workbook(source)
    ws = wb[wb.sheetnames[0]]

    written = 0
    for _, row in universe.iterrows():
        cell = f"{shares_column}{int(row['sheet_row'])}"
        if row["is_duplicate_row"]:
            ws[cell] = 0                      # the second HST line
            continue
        shares = int(table["shares"].get(row["permno"], 0))
        ws[cell] = shares
        if shares > 0:
            written += 1

    ws[f"{shares_column}{SHEET_RF_ROW}"] = float(riskfree_amount)

    # The template's subtotals read =SUM(x22:x524) but the firm block runs to
    # row 526, so as delivered they silently omit the last two firms (ZBRA and
    # ZTS). All three parts of the assignment use the same sheet, so all three
    # subtotal cells are repaired, not just the one we fill in.
    old_formulas, new_formulas = {}, {}
    for cell, column in [("I16", "I"), ("L16", "L"), ("O16", "O")]:
        old_formulas[cell] = ws[cell].value
        ws[cell] = f"=SUM({column}{SHEET_FIRST_FIRM_ROW}:{column}{SHEET_LAST_FIRM_ROW})"
        new_formulas[cell] = ws[cell].value

    if group:
        for cell, key in [("C2", "group"), ("C3", "student1"), ("C4", "student1_nr"),
                          ("C5", "student2"), ("C6", "student2_nr"),
                          ("C7", "student3"), ("C8", "student3_nr")]:
            if key in group:
                ws[cell] = group[key]

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)

    return {
        "rows_with_shares": written,
        "riskfree_cell": f"{shares_column}{SHEET_RF_ROW}",
        "old_subtotal_formulas": old_formulas,
        "new_subtotal_formulas": new_formulas,
        "path": target,
    }
